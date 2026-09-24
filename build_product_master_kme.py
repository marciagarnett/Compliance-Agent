"""
One-time script that expands data/product_master.csv further, adding HP's
non-Poly product lines (notebooks, desktops, workstations, thin clients,
retail/RPOS systems, and displays) from ../HP Products under KME .xlsx.

Not part of the running application - kept for transparency/reproducibility,
same reason build_data.py and build_product_master.py are kept. Idempotent:
anything whose sku is already in product_master.csv is left alone.

Source and what's in it (see source_index.md for the full write-up): most
sheets in this workbook are KME's own documentation-*program* tracking -
project codenames, ODM contacts, doc-kit/QSG part numbers, code-freeze and
GA dates - not a finished-goods parts catalog the way the ASCM export was
for Poly (see build_product_master.py). Concretely, that means:

- Most sheets have no real per-unit part number at all. Where a "Model
  Number" column exists, it is very often a WILDCARDED FAMILY PATTERN (e.g.
  "24-cr0xxx", "GT16-0xxx, GT16-000") covering many actual configurations,
  not one orderable SKU.
- Per an explicit decision on this project, `sku` is populated from that
  Model Number column as-is (wildcards included) when the sheet has one.
  When a sheet has no Model Number column at all (Commercial Solutions-
  Mobile Wks/TC, Commercial - Retail Mobility, Commercial Solutions-CMIT-
  RPOS, Displays - RPOS IQ), this script falls back to a "PROJ-<Project
  Code>" identifier instead - an internal KME project code, not a customer-
  facing part number or family pattern either. This is a materially weaker
  identifier than either of the above and was not separately confirmed the
  way the Model Number fallback was - flagged here and in the project
  write-up so it can be revisited.
- Several sheets have a column literally named "RMN" (or "Regulatory
  Model") whose values repeat identically across clearly unrelated products
  (e.g. the same value on an All-in-One desktop, an Envy, and two different
  OMEN gaming towers) - not what this app's rmn field means anywhere else
  (one distinct regulatory model number per product). Per an explicit
  decision on this project, it is still populated as-is; RMN lookup on a
  shared value will resolve to whichever product this script listed first.
- One sheet, "AMO WOR's without kits", genuinely is a parts list with real,
  specific part numbers (e.g. "L10283-005") - handled on its own below, and
  rows marked STATUS == "CANCELLED" are skipped as never-shipped.
- The "Poly" sheet is excluded entirely (handled/excluded already - see
  source_index.md - it is an unreleased-product roadmap, not a catalog).
- All five Display sheets are collapsed into one "Display" category for
  now; they have no compliance data to differentiate yet regardless of
  sub-line (business/office, workstation, consumer, gaming, RPOS), so
  splitting further can wait until it would actually matter.

Column names vary by sheet (e.g. "Model Name(s)" vs "Model Descriptions",
"RMN" vs "Regulatory Model", "Product Sub-Type" vs "sub-type") - resolved via
the alias sets below rather than fixed column positions, since header rows
also start at different row numbers per sheet.
"""
import csv
import re
import openpyxl

SRC = "../HP Products under KME .xlsx"
PM_PATH = "data/product_master.csv"

MODEL_NAME_ALIASES = {"model name(s)", "model names", "model descriptions"}
MODEL_NUMBER_ALIASES = {"model number", "model number (s)", "model numbers"}
RMN_ALIASES = {"rmn", "regulatory model"}
SUBTYPE_ALIASES = {"product sub-type", "product sub type", "sub-type", "subtype"}
PROJECT_CODE_ALIASES = {"project\ncode", "project code"}
PROGRAM_ALIASES = {"program"}
PROJECT_NAME_ALIASES = {"project name"}

# name -> (header_row, category)
SHEETS = {
    "Cons Notebooks": (1, "Consumer Notebook"),
    "Cons Desktop": (1, "Consumer Desktop"),
    "Commercial Notebook": (1, "Commercial Notebook"),
    "Commercial Solutions-Mobile Wks": (1, "Commercial Mobile Workstation"),
    "Commercial Solutions-Mobile TC": (1, "Commercial Mobile Thin Client"),
    "Commercial - Retail Mobility": (1, "Commercial Retail Mobility"),
    "CM Desktop": (1, "Commercial Desktop"),
    "Commercial Solutions-CMIT Wkstn": (1, "Commercial Workstation"),
    "Commercial Solutions-CMIT TC": (1, "Commercial Thin Client"),
    "Commercial Solutions-CMIT-RPOS": (2, "Commercial RPOS"),
    "Displays - CM  BO": (3, "Display"),
    "Displays - CM Workstation  TB": (3, "Display"),
    "Displays - CN   2G": (7, "Display"),
    "Displays - CN Gaming  2H": (7, "Display"),
    "Displays - RPOS  IQ": (2, "Display"),
}


def norm(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def col_map(header_row):
    m = {}
    for i, h in enumerate(header_row):
        n = norm(h)
        if n:
            m.setdefault(n, i)
    return m


def find_col(cmap, aliases):
    for a in aliases:
        if a in cmap:
            return cmap[a]
    return None


def get(row, idx):
    if idx is None or idx >= len(row):
        return None
    v = row[idx]
    if v is None:
        return None
    return str(v).strip() or None


def extract_program_sheets(wb):
    rows_out = []
    for sheet_name, (header_row, category) in SHEETS.items():
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=True)
        header = None
        for i, row in enumerate(rows_iter, start=1):
            if i == header_row:
                header = row
                break
        cmap = col_map(header)
        name_idx = find_col(cmap, MODEL_NAME_ALIASES)
        num_idx = find_col(cmap, MODEL_NUMBER_ALIASES)
        rmn_idx = find_col(cmap, RMN_ALIASES)
        sub_idx = find_col(cmap, SUBTYPE_ALIASES)
        proj_code_idx = find_col(cmap, PROJECT_CODE_ALIASES)
        program_idx = find_col(cmap, PROGRAM_ALIASES)
        proj_name_idx = find_col(cmap, PROJECT_NAME_ALIASES)

        for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row,
                                 max_col=ws.max_column, values_only=True):
            model_name = get(row, name_idx)
            model_number = get(row, num_idx)
            if not model_name and not model_number:
                continue  # blank spacer row or nothing usable

            product_name = model_name or model_number or get(row, program_idx) or get(row, proj_name_idx)
            if not product_name:
                continue

            if model_number:
                sku = model_number
            else:
                proj_code = get(row, proj_code_idx)
                sku = f"PROJ-{proj_code}" if proj_code else None
            if not sku:
                continue

            rmn = get(row, rmn_idx) or ""
            cat = category
            if sub_idx is not None:
                sub = get(row, sub_idx)
                # Keep the sheet-level category (Consumer Notebook, etc.) rather
                # than overwriting with sub-type text - sub-type here is more
                # like a further breakdown (e.g. "DSK - Other") than a clean
                # category name, and 32 categories is already a lot to add
                # compliance data for.
                _ = sub

            rows_out.append({
                "product_name": product_name,
                "sku": sku,
                "rmn": rmn,
                "category": cat,
            })
    return rows_out


def extract_amo_parts(wb):
    ws = wb["AMO WOR's without kits"]
    header = next(ws.iter_rows(min_row=1, max_row=1, max_col=ws.max_column, values_only=True))
    cmap = col_map(header)
    status_idx = cmap.get("status")
    subtype_idx = cmap.get("sub-type")
    product_idx = cmap.get("product")
    pn_idx = cmap.get("part numer") or cmap.get("part number")
    desc_idx = cmap.get("description")

    rows_out = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column, values_only=True):
        status = (get(row, status_idx) or "").upper()
        if status == "CANCELLED":
            continue
        pn = get(row, pn_idx)
        if not pn:
            continue
        desc = get(row, desc_idx) or get(row, product_idx) or pn
        category = get(row, subtype_idx) or "Accessory / Service Part"
        rows_out.append({
            "product_name": desc,
            "sku": pn,
            "rmn": "",
            "category": category,
        })
    return rows_out


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

    with open(PM_PATH, newline="", encoding="utf-8") as f:
        existing = list(csv.DictReader(f))
    existing_skus = {row["sku"].strip() for row in existing if row["sku"]}

    candidates = extract_program_sheets(wb) + extract_amo_parts(wb)

    new_rows = []
    seen_skus = set()
    skipped_existing = 0
    skipped_dupe = 0
    for r in candidates:
        sku = r["sku"].strip()
        if sku in existing_skus:
            skipped_existing += 1
            continue
        if sku in seen_skus:
            skipped_dupe += 1
            continue
        seen_skus.add(sku)
        new_rows.append(r)

    new_rows.sort(key=lambda r: (r["category"], r["product_name"]))

    print(f"existing rows kept as-is: {len(existing)}")
    print(f"candidates extracted: {len(candidates)}")
    print(f"skipped (already in product_master.csv): {skipped_existing}")
    print(f"skipped (duplicate sku within this run): {skipped_dupe}")
    print(f"new rows added: {len(new_rows)}")

    with open(PM_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["product_name", "sku", "rmn", "category"])
        writer.writeheader()
        writer.writerows(existing)
        writer.writerows(new_rows)


if __name__ == "__main__":
    main()
